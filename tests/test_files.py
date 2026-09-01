"""Unit tests for the file-upload blob store + text extraction (no DB)."""

import hashlib
from io import BytesIO

import pytest

from cortex import files


@pytest.fixture(autouse=True)
def _tmp_files_dir(tmp_path, monkeypatch):
    monkeypatch.setenv("CORTEX_FILES_DIR", str(tmp_path))
    yield tmp_path


def test_save_blob_is_content_addressed():
    content = b"hello cortex"
    sha, path, size = files.save_blob(content)
    assert sha == hashlib.sha256(content).hexdigest()
    assert size == len(content)
    assert files.read_blob(path) == content


def test_save_blob_dedupes_identical_content():
    content = b"same bytes, uploaded twice"
    _, path1, _ = files.save_blob(content)
    _, path2, _ = files.save_blob(content)
    assert path1 == path2  # second upload didn't write a new blob


def test_save_blob_different_content_different_path():
    _, path1, _ = files.save_blob(b"content A")
    _, path2, _ = files.save_blob(b"content B")
    assert path1 != path2


def test_extract_text_plain_and_markdown():
    assert files.extract_text("notes.txt", "text/plain", b"plain text body") == "plain text body"
    assert files.extract_text("readme.md", "text/markdown", b"# heading\n\nbody") == "# heading\n\nbody"


def test_extract_text_unsupported_type_returns_empty_not_raises():
    assert files.extract_text("photo.jpg", "image/jpeg", b"\xff\xd8\xff\xe0binarydata") == ""


def test_extract_text_corrupt_pdf_returns_empty_not_raises():
    assert files.extract_text("broken.pdf", "application/pdf", b"not a real pdf") == ""


def test_extract_text_docx_round_trip():
    from docx import Document

    doc = Document()
    doc.add_paragraph("first paragraph")
    doc.add_paragraph("second paragraph")
    buf = BytesIO()
    doc.save(buf)
    text = files.extract_text("report.docx", "", buf.getvalue())
    assert "first paragraph" in text
    assert "second paragraph" in text
